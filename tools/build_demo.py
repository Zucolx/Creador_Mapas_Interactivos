#!/usr/bin/env python3
"""
Genera demo/demo-data.js: el proyecto de ejemplo (mapas L1 y L2) con todas
las imagenes incrustadas como data-URL, listo para abrirse en el editor.

Uso:
  python3 tools/build_demo.py \
      --pptx  "1. LAYOUT_QUIMICOS_ENVASE_L1_11-2025.pptx" \
      --sga   "ETIQUETAS_FORMATO_SGA.pdf" \
      --plano-l1 plano_L1.webp --plano-l2 plano_L2.webp \
      --out demo/demo-data.js

Requiere: python-pptx, Pillow y poppler-utils (pdftoppm / pdftotext).
"""
import argparse, base64, io, json, re, subprocess, tempfile, os
from PIL import Image
from pptx import Presentation

# ---------------------------------------------------------------------------
# Biblioteca de quimicos. Estilo del pictograma = relleno + aro (como en el
# PowerPoint original). "fichas" = paginas del PDF de etiquetas SGA.
# ---------------------------------------------------------------------------
CHEMICALS = [
    # id,             nombre,                       relleno,   aro,       paginas SGA
    ("topax32",       "TOPAX 32",                   "#FFFFFF", "#000000", [18]),
    ("topax58",       "TOPAX 58",                   "#FFFFFF", "#FF0000", [19]),
    ("topax66",       "TOPAX 66",                   "#FFFFFF", "#FFD800", [16, 17]),
    ("oxonia",        "OXONIA ACTIVE",              "#FFFFFF", "#00B0F0", [14, 15]),
    ("goma",          "GOMA (Moplast 10-14)",       "#FFFFFF", "#00B050", [13]),
    ("rimazon_oxo",   "RIMAZON OXO",                "#0070C0", "#000000", [22]),
    ("rimazon_sauer", "RIMAZON SAUER",              "#FF0000", "#000000", [21]),
    ("dicolube",      "DICOLUBE CR",                "#00B050", "#000000", [1]),
    ("dryexx",        "DRYEXX GF",                  "#00B050", "#000000", []),
    ("su221",         "SU 221",                     "#0070C0", "#00B050", [2]),
    ("stabilon",      "STABILON ACI",               "#0070C0", "#00B050", [20]),
    ("soda",          "SODA CÁUSTICA",              "#7030A0", "#7030A0", [5, 6, 7, 8, 9, 10, 11]),
    ("acido_nitrico", "ÁCIDO NÍTRICO",              "#000000", "#FF0000", [23, 24, 25, 26, 27]),
    ("alcohol",       "ALCOHOL",                    "#00B0F0", "#00B0F0", []),
    ("dcitrus",       "D-CITRUS",                   "#FFFF00", "#000000", [3]),
    # Presentes en el PDF de etiquetas pero sin ubicacion en los mapas actuales
    ("amoniaco",      "AMONIACO",                   "#FFFFFF", "#7F7F7F", [4]),
    ("fibra_soda",    "FIBRA HÚMEDA CON RESIDUOS DE SODA", "#C9A3E6", "#7030A0", [12]),
    ("sc2000",        "SC-2000",                    "#FFFFFF", "#ED7D31", [28]),
    ("divosan",       "DIVOSAN FORTE",              "#ED7D31", "#000000", [29]),
    ("sol_3601q",     "16-3601Q Solución de limpieza", "#A5A5A5", "#000000", [31]),
    ("v483a",         "V483A-C Tinta de impresión", "#404040", "#404040", [32]),
    ("v483ab",        "V483AB-C Tinta de impresión", "#404040", "#FFC000", [33]),
    ("v721d",         "V721-D Solvente",            "#FFFFFF", "#404040", [34]),
]

# ---------------------------------------------------------------------------
# Mapas. Coordenadas en pixeles del plano original (2000 x 1415). Pueden ser
# negativas: el plano se amplia con margen blanco para ubicar zonas externas
# (tanques, barrilera, cuartos), igual que en el layout del PowerPoint.
# area: (id, nombre, descripcion, rect x0,y0,x1,y1, [diapositivas de fotos],
#        [(quimico, x, y), ...])
# ---------------------------------------------------------------------------
L1 = dict(
    id="L1", name="Línea 1", title="Mapa de sustancias químicas", subtitle="Línea 1 · Envase / Barrilera",
    plano="plano_l1",
    areas=[
        ("tanques_soda", "Tanques de soda cáustica", "Tanques de almacenamiento de soda cáustica (exterior).",
         (690, -150, 990, -20), [6], [("soda", 790, -85), ("soda", 890, -85)]),
        ("dioxido_cloro", "Generador de dióxido de cloro", "Precursores Rimazon Oxo y Rimazon Sauer.",
         (1280, -150, 1480, -20), [5], [("rimazon_sauer", 1350, -85), ("rimazon_oxo", 1410, -85)]),
        ("barrilera", "Barrilera", "Área de barrilera.",
         (1660, -150, 1990, -20), [9],
         [("acido_nitrico", 1720, -85), ("topax66", 1780, -85), ("acido_nitrico", 1870, -85), ("soda", 1930, -85)]),
        ("lavadora", "Lavadora de botellas", "Dosificación de soda cáustica en la lavadora.",
         (405, 160, 965, 355), [7],
         [("soda", 560, 318), ("soda", 640, 318), ("soda", 720, 318), ("soda", 800, 318)]),
        ("estacion_limpieza", "Estación de limpieza (Topax / Oxonia)", "Estación de dosificación de productos de limpieza y desinfección.",
         (255, 820, 380, 945), [10],
         [("topax66", 295, 860), ("topax58", 340, 860), ("topax32", 295, 905), ("oxonia", 340, 905)]),
        ("tanques_cip", "Tanques CIP (soda / ácido nítrico)", "Tanques de soda y ácido nítrico del sistema CIP.",
         (680, 870, 960, 1035), [8],
         [("acido_nitrico", 735, 915), ("soda", 790, 915), ("soda", 905, 945), ("acido_nitrico", 905, 995)]),
        ("cuarto_mp_l1", "Cuarto de materia prima L1", "Almacenamiento de productos químicos de la línea 1.",
         (1360, 1085, 1620, 1230), [4],
         [("topax66", 1405, 1130), ("topax32", 1450, 1130), ("acido_nitrico", 1495, 1130),
          ("dcitrus", 1545, 1130), ("goma", 1590, 1130), ("oxonia", 1405, 1180), ("topax58", 1450, 1180)]),
        ("almacen_lubricante", "Almacén Dicolube CR / SU 221", "Lubricante de cadenas y aditivo para lavado de botellas.",
         (1700, 1100, 1965, 1290), [3], [("dicolube", 1880, 1160), ("su221", 1880, 1230)]),
    ],
    # Pictogramas sin area (abren la ficha directamente)
    markers=[("goma", 1230, 700), ("goma", 1385, 745), ("alcohol", 600, 1300)],
    labels=[],
)

L2 = dict(
    id="L2", name="Línea 2", title="Mapa de sustancias químicas", subtitle="Línea 2 · Envase",
    plano="plano_l2",
    areas=[
        ("tanques_cip", "Tanques CIP (soda / ácido nítrico)", "Tanques de soda cáustica y ácido nítrico.",
         (60, 95, 340, 200), [], [("acido_nitrico", 110, 148), ("soda", 170, 148), ("soda", 230, 148), ("acido_nitrico", 290, 148)]),
        ("cuarto_lubricante", "Cuarto de lubricante", "Lubricante de cadenas Dryexx GF.",
         (-215, 430, -15, 650), [], [("dryexx", -115, 560)]),
        ("estacion_limpieza", "Estación de limpieza (Topax / Oxonia)", "Estación de dosificación de productos de limpieza y desinfección.",
         (70, 565, 155, 665), [], [("topax66", 92, 595), ("topax58", 132, 595), ("topax32", 92, 635), ("oxonia", 132, 635)]),
        ("punto_limpieza_llenadora", "Punto de limpieza llenadora", "Topax 66 y alcohol para limpieza de la llenadora.",
         (60, 915, 185, 1040), [], [("topax66", 122, 955), ("alcohol", 122, 1000)]),
        ("etiquetadora", "Etiquetadora", "Goma para etiquetas.",
         (-5, 1045, 175, 1205), [], [("goma", 155, 1065), ("goma", 155, 1185)]),
        ("dioxido_cloro", "Generador de dióxido de cloro", "Precursores Rimazon Oxo y Rimazon Sauer.",
         (480, 790, 595, 890), [], [("rimazon_oxo", 538, 822), ("rimazon_sauer", 538, 860)]),
        ("lubricante_cadenas", "Lubricante de cadenas", "Dosificación de Dryexx GF.",
         (1040, 912, 1185, 992), [], [("dryexx", 1112, 952)]),
        ("dosificacion_lavadora", "Dosificación lavadora", "Stabilon, soda cáustica y Topax 66.",
         (465, 235, 645, 330), [], [("stabilon", 510, 282), ("soda", 555, 282), ("topax66", 600, 282)]),
        ("cuarto_mp_l2", "Cuarto de materia prima L2", "Almacenamiento de productos químicos de la línea 2.",
         (925, 1320, 1075, 1500), [],
         [("goma", 980, 1345), ("acido_nitrico", 980, 1380), ("rimazon_sauer", 1020, 1380),
          ("oxonia", 980, 1415), ("topax32", 1020, 1415), ("topax58", 980, 1450), ("topax66", 1020, 1450),
          ("dcitrus", 980, 1485), ("rimazon_oxo", 1020, 1485)]),
    ],
    markers=[("oxonia", 15, 760), ("soda", 800, 455), ("soda", 895, 455), ("soda", 990, 455), ("soda", 1085, 455)],
    labels=[("USTED ESTÁ AQUÍ", 860, 1540)],
)

MARGIN = 45


def data_url(img, fmt="WEBP", quality=72, lossless=False):
    buf = io.BytesIO()
    if fmt == "WEBP":
        img.save(buf, "WEBP", quality=quality, method=6, lossless=lossless)
        mime = "image/webp"
    else:
        img.save(buf, "PNG", optimize=True)
        mime = "image/png"
    return f"data:{mime};base64," + base64.b64encode(buf.getvalue()).decode()


def fit(img, max_side):
    w, h = img.size
    s = min(1.0, max_side / max(w, h))
    return img.resize((round(w * s), round(h * s)), Image.LANCZOS) if s < 1 else img


# ----------------------------- fichas SGA ----------------------------------
def parse_label(text):
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    signal = next((l for l in lines[:6] if l.upper() in ("PELIGRO", "ATENCIÓN", "ATENCION")), "")
    rest = [l for l in lines if l != signal]
    title = rest[0] if rest else ""
    subtitle = rest[1] if len(rest) > 1 and not rest[1].startswith("Componentes") else ""
    hazards = [re.sub(r"\s+", " ", l) for l in lines if re.match(r"^H\d{3}", l)]
    volume = next((l.split(":", 1)[1].strip() for l in lines if l.upper().startswith("VOLUMEN")), "")
    m = re.match(r"^([\d.,]+\s*(kg|litros?|gal|ml|m3|l)\b)", volume, re.I)
    volume = m.group(1) if m else volume
    emergency = ""
    for i, l in enumerate(lines):
        if l.upper().startswith("CONTACTO DE EMERGENCIA"):
            emergency = l.split(":", 1)[1].strip(" .") if ":" in l else ""
            if not emergency and i + 1 < len(lines):
                emergency = lines[i + 1]
            break
    return dict(title=title, subtitle=subtitle, signal=signal.upper().replace("ATENCION", "ATENCIÓN"),
                hazards=hazards, volume=volume, emergency=emergency)


def build_fichas(pdf, assets):
    tmp = tempfile.mkdtemp()
    subprocess.run(["pdftoppm", "-r", "150", "-png", pdf, os.path.join(tmp, "p")], check=True)
    files = sorted(f for f in os.listdir(tmp) if f.endswith(".png"))
    by_page = {}
    for f in files:
        n = int(re.search(r"(\d+)\.png$", f).group(1))
        text = subprocess.run(["pdftotext", "-raw", "-f", str(n), "-l", str(n), pdf, "-"],
                              capture_output=True, text=True).stdout
        meta = parse_label(text)
        img = fit(Image.open(os.path.join(tmp, f)).convert("RGB"), 1400)
        aid = f"sga_p{n:02d}"
        assets[aid] = data_url(img, quality=72)
        cap = " · ".join(x for x in (meta["subtitle"], meta["volume"]) if x)
        by_page[n] = dict(img=aid, caption=cap or meta["title"], signal=meta["signal"],
                          hazards=meta["hazards"], emergency=meta["emergency"], source=f"Etiquetas SGA, pág. {n}")
    return by_page


# ----------------------------- fotos PPTX ----------------------------------
def slide_photos(prs, slide_no):
    """Imagenes de una diapositiva (recortadas como en la presentacion)."""
    out = []
    for sh in prs.slides[slide_no - 1].shapes:
        if sh.shape_type != 13:
            continue
        try:
            blob = sh.image.blob
        except Exception:  # iconos SVG enlazados
            continue
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        w, h = img.size
        l, t, r, b = sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom
        box = (max(0, round(l * w)), max(0, round(t * h)), min(w, round(w - r * w)), min(h, round(h - b * h)))
        if box[2] - box[0] > 50 and box[3] - box[1] > 50:
            img = img.crop(box)
        if min(img.size) < 120:
            continue
        out.append((sh.left, img))
    out.sort(key=lambda x: x[0])
    return [img for _, img in out]


# ----------------------------- mapas ---------------------------------------
def build_map(cfg, plano_path, prs, assets, used):
    base = Image.open(plano_path).convert("RGB")
    W, H = base.size
    xs, ys = [0, W], [0, H]
    for a in cfg["areas"]:
        x0, y0, x1, y1 = a[3]
        xs += [x0, x1]; ys += [y0, y1]
        for _, x, y in a[5]:
            xs += [x - 20, x + 20]; ys += [y - 20, y + 20]
    for _, x, y in cfg["markers"]:
        xs += [x - 20, x + 20]; ys += [y - 20, y + 20]
    for _, x, y in cfg["labels"]:
        xs += [x - 150, x + 150]; ys += [y - 30, y + 30]
    ox = -min(xs) + MARGIN if min(xs) < 0 else 0
    oy = -min(ys) + MARGIN if min(ys) < 0 else 0
    NW = max(max(xs) + ox + (MARGIN if max(xs) > W else 0), W + ox)
    NH = max(max(ys) + oy + (MARGIN if max(ys) > H else 0), H + oy)
    canvas = Image.new("RGB", (int(NW), int(NH)), "white")
    canvas.paste(base, (int(ox), int(oy)))
    canvas = canvas.convert("L")
    aid = cfg["plano"]
    assets[aid] = data_url(canvas, quality=80)
    fx = lambda x: round((x + ox) / NW * 100, 3)
    fy = lambda y: round((y + oy) / NH * 100, 3)

    areas, markers = [], []
    mid = 0
    for (a_id, name, desc, (x0, y0, x1, y1), slides, items) in cfg["areas"]:
        photos = []
        for s in slides:
            for k, img in enumerate(slide_photos(prs, s)):
                pid = f"{cfg['id']}_{a_id}_{k}"
                assets[pid] = data_url(fit(img, 1600), quality=68)
                photos.append(dict(img=pid, pins=[]))
        chems = []
        for c, x, y in items:
            mid += 1
            markers.append(dict(id=f"m{mid}", chem=c, x=fx(x), y=fy(y), area=a_id))
            if c not in chems:
                chems.append(c)
            used.add(c)
        areas.append(dict(id=a_id, name=name, desc=desc, x=fx(x0), y=fy(y0),
                          w=round((x1 - x0) / NW * 100, 3), h=round((y1 - y0) / NH * 100, 3),
                          photos=photos, chems=chems))
    for c, x, y in cfg["markers"]:
        mid += 1
        markers.append(dict(id=f"m{mid}", chem=c, x=fx(x), y=fy(y), area=None))
        used.add(c)
    labels = [dict(id=f"t{i}", text=t, x=fx(x), y=fy(y), size=1.0) for i, (t, x, y) in enumerate(cfg["labels"], 1)]
    return dict(id=cfg["id"], name=cfg["name"], title=cfg["title"], subtitle=cfg["subtitle"],
                background=aid, bgWidth=int(NW), bgHeight=int(NH), markerSize=1.0,
                areas=areas, markers=markers, labels=labels)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--sga", required=True)
    ap.add_argument("--plano-l1", required=True)
    ap.add_argument("--plano-l2", required=True)
    ap.add_argument("--out", default="demo/demo-data.js")
    a = ap.parse_args()

    assets = {}
    prs = Presentation(a.pptx)
    fichas = build_fichas(a.sga, assets)
    used = set()
    maps = [build_map(L1, a.plano_l1, prs, assets, used), build_map(L2, a.plano_l2, prs, assets, used)]
    chems = []
    for cid, name, fill, ring, pages in CHEMICALS:
        chems.append(dict(id=cid, name=name, fill=fill, ring=ring,
                          fichas=[fichas[p] for p in pages if p in fichas]))
    project = dict(app="creador-mapas-interactivos", version=1, name="Mapas de sustancias químicas · Envase",
                   chemicals=chems, maps=maps, assets=assets)
    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as f:
        f.write("/* Proyecto de ejemplo generado por tools/build_demo.py */\n")
        f.write("window.DEMO_PROJECT = ")
        json.dump(project, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    size = os.path.getsize(a.out) / 1e6
    print(f"{a.out}: {size:.1f} MB, {len(assets)} imágenes, {len(chems)} químicos")
    for m in maps:
        print(f"  {m['id']}: {len(m['areas'])} áreas, {len(m['markers'])} pictogramas, plano {m['bgWidth']}x{m['bgHeight']}")


if __name__ == "__main__":
    main()
